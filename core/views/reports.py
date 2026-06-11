from django.shortcuts import render, get_object_or_404
from django.utils import timezone
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.views import View

from weasyprint import HTML
import openpyxl

from users.mixins import CounselorRequiredMixin, AdminRequiredMixin
from core.models import Student, Meeting
from core.services import _build_annual_report_data

class AnnualReportView(AdminRequiredMixin, View):
    def get(self, request, *args, **kwargs):
        year = self._get_year(request)
        data = _build_annual_report_data(year)
        data['available_years'] = self._available_years()
        return render(request, 'core/annual_report.html', data)

    def _get_year(self, request):
        year_param = request.GET.get('year', '').strip()
        if year_param.isdigit():
            return int(year_param)
        return timezone.now().year

    def _available_years(self):
        years = set()
        for m in Meeting.objects.all():
            years.add(m.date_time.year)
        years.add(timezone.now().year)
        return sorted(years, reverse=True)


class AnnualReportPDFView(AdminRequiredMixin, View):
    def get(self, request, *args, **kwargs):
        year_param = request.GET.get('year', '').strip()
        year = int(year_param) if year_param.isdigit() else timezone.now().year

        data = _build_annual_report_data(year)
        data['today'] = timezone.now().date()

        html_string = render_to_string('core/annual_report_pdf.html', data)
        pdf = HTML(string=html_string).write_pdf()

        response = HttpResponse(pdf, content_type='application/pdf')
        response['Content-Disposition'] = f'inline; filename="godisnji_izvjestaj_{year}.pdf"'
        return response


class AnnualReportExcelView(AdminRequiredMixin, View):
    def get(self, request, *args, **kwargs):
        year_param = request.GET.get('year', '').strip()
        year = int(year_param) if year_param.isdigit() else timezone.now().year

        data = _build_annual_report_data(year)

        wb = openpyxl.Workbook()

        # List 1: Sažetak
        ws = wb.active
        ws.title = 'Sažetak'
        ws.append([f'Godišnji izvještaj AAC-a za {year}. godinu'])
        ws.append([])
        ws.append(['Aktivni studenti', data['active_students']])
        ws.append(['Ukupno sastanaka', data['total_meetings']])
        ws.append(['Sati vršnjačke podrške', data['peer_support_hours']])

        # List 2: Po savjetniku
        ws2 = wb.create_sheet('Po savjetniku')
        ws2.append(['Savjetnik', 'Broj sastanaka', 'Broj studenata'])
        for row in data['per_counselor']:
            ws2.append([row['counselor'].full_name, row['meeting_count'], row['student_count']])

        # List 3: Vrste teškoća
        ws3 = wb.create_sheet('Vrste teškoća')
        ws3.append(['Vrsta teškoće', 'Broj studenata'])
        for row in data['per_disability']:
            ws3.append([row['disability'].name, row['count']])

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="godisnji_izvjestaj_{year}.xlsx"'
        wb.save(response)
        return response
    
class StudentReportPDFView(CounselorRequiredMixin, View):
    def get(self, request, *args, **kwargs):
        student = get_object_or_404(
            Student.objects.filter(counselors=request.user.counselor_profile),
            pk=kwargs['pk']
        )

        context = {
            'student': student,
            'meetings': student.meetings.filter(is_active=True).order_by('-date_time'),
            'accommodations': student.accommodations.all().order_by('-start_date'),
            'documents': student.documents.all().order_by('-uploaded_at'),
            'counselor': request.user.counselor_profile,
            'today': timezone.now().date(),
        }

        html_string = render_to_string('core/student_report_pdf.html', context)
        pdf = HTML(string=html_string).write_pdf()

        response = HttpResponse(pdf, content_type='application/pdf')
        filename = f"izvjestaj_student_{student.last_name}_{student.first_name}.pdf"
        response['Content-Disposition'] = f'inline; filename="{filename}"'
        return response